"""会議録 docx から、段落の文字を「抜けなく」読み出す部品。

python-docx の p.text は、段落の直下の文字（w:r の w:t）だけを読むため、次の文字が消える：
  - 均等割付の EQ フィールド（instrText の中の eq \\o\\ad(委員長,　　　)）
  - ルビ（w:ruby の中の親文字）
  - フィールドの結果、ハイパーリンク・変更履歴（w:ins）の中の文字
  - 表（w:tbl）の中の段落（doc.paragraphs に入らない）
ここでは文書の本文を上から順に歩き、表の中の段落も含めて全部の段落を返す。
AI は使わない。文字は足しも削りもしない（EQ だけは見た目の文字を取り出す）。
"""
import re

import docx
from docx.oxml.ns import qn

W_P, W_TBL, W_R, W_T = qn("w:p"), qn("w:tbl"), qn("w:r"), qn("w:t")
W_TAB, W_BR, W_CR = qn("w:tab"), qn("w:br"), qn("w:cr")
W_FLDCHAR, W_INSTR, W_FLDSIMPLE = qn("w:fldChar"), qn("w:instrText"), qn("w:fldSimple")
W_RUBY, W_RT, W_RUBYBASE = qn("w:ruby"), qn("w:rt"), qn("w:rubyBase")
W_DEL, W_DELTEXT = qn("w:del"), qn("w:delText")
W_TXBX, W_DRAWING, W_PICT = qn("w:txbxContent"), qn("w:drawing"), qn("w:pict")
W_TR, W_TC, W_SDT, W_SDTCONTENT = qn("w:tr"), qn("w:tc"), qn("w:sdt"), qn("w:sdtContent")
W_FLDTYPE = qn("w:fldCharType")
W_PPR, W_RPR = qn("w:pPr"), qn("w:rPr")
W_INSTR_ATTR = qn("w:instr")

_EQ_CMD = re.compile(r"\\[a-zA-Z]+(?:\s*-?\d+)?")


def eq_visible_text(instr):
    """EQ フィールドの命令文から、画面に見える文字を取り出す。

    例: 'eq \\o\\ad(委員長,　　　　　)' → '委員長'
        'EQ \\o\\ad(\\s\\up 0(出席),\\s\\up 0(　　))' → '出席'
    \\o（重ね書き）の引数は「文字」と「幅をとるための空白」なので、最初の引数を使う。
    """
    s = instr.strip()
    if not re.match(r"(?i)^eq\b", s):
        return None
    body = s[2:]
    start = body.find("(")
    if start < 0:
        return _EQ_CMD.sub("", body).strip()
    depth, args, cur = 0, [], []
    for ch in body[start:]:
        if ch == "(":
            depth += 1
            if depth == 1:
                continue
        elif ch == ")":
            depth -= 1
            if depth == 0:
                args.append("".join(cur))
                break
        elif ch == "," and depth == 1:
            args.append("".join(cur))
            cur = []
            continue
        cur.append(ch)
    first = args[0] if args else ""
    first = _EQ_CMD.sub("", first)
    return first.replace("(", "").replace(")", "").strip(" ")


class _ParaReader:
    def __init__(self):
        self.out = []
        self.fields = []   # スタック: {"instr": str, "phase": "instr"|"result", "result": str}
        self.notes = {"eq": [], "ruby": [], "field_other": [], "textbox": 0}

    def emit(self, s):
        if any(f["phase"] == "instr" for f in self.fields):
            return  # 命令文の途中の文字は見えない
        self.out.append(s)
        if self.fields:
            self.fields[-1]["result"] += s

    def walk(self, el):
        for ch in el:
            tag = ch.tag
            if tag == W_T:
                self.emit(ch.text or "")
            elif tag == W_TAB:
                self.emit("\t")
            elif tag in (W_BR, W_CR):
                self.emit("\n")
            elif tag == W_FLDCHAR:
                t = ch.get(W_FLDTYPE)
                if t == "begin":
                    self.fields.append({"instr": "", "phase": "instr", "result": ""})
                elif t == "separate" and self.fields:
                    self.fields[-1]["phase"] = "result"
                elif t == "end" and self.fields:
                    self._close(self.fields.pop())
            elif tag == W_INSTR:
                if self.fields:
                    self.fields[-1]["instr"] += ch.text or ""
            elif tag == W_FLDSIMPLE:
                self.fields.append({"instr": ch.get(W_INSTR_ATTR) or "",
                                    "phase": "result", "result": ""})
                self.walk(ch)
                self._close(self.fields.pop())
            elif tag == W_RUBY:
                base = ch.find(W_RUBYBASE)
                rt = ch.find(W_RT)
                before = len("".join(self.out))
                if base is not None:
                    self.walk(base)
                base_text = "".join(self.out)[before:]
                rt_text = "".join(t.text or "" for t in rt.iter(W_T)) if rt is not None else ""
                self.notes["ruby"].append({"base": base_text, "ruby": rt_text})
            elif tag in (W_PPR, W_RPR):
                continue  # 書式の設定（タブ位置の w:tab などは文字ではない）
            elif tag in (W_DEL, W_DELTEXT):
                continue  # 変更履歴で消された文字は読まない
            elif tag in (W_DRAWING, W_PICT, W_TXBX):
                if any(True for _ in ch.iter(W_TXBX)) or tag == W_TXBX:
                    self.notes["textbox"] += 1
                continue  # テキストボックスの中は段落の外の文字なので別扱い
            else:
                self.walk(ch)  # w:r, w:hyperlink, w:ins, w:smartTag, w:sdt など

    def _close(self, f):
        instr = f["instr"].strip()
        if re.match(r"(?i)^eq\b", instr):
            vis = eq_visible_text(instr)
            if not f["result"] and vis:
                # 結果が空の EQ は、命令文の中の文字が見た目の文字
                self.emit(vis)
            self.notes["eq"].append({"instr": instr, "visible": vis or f["result"]})
        elif instr:
            self.notes["field_other"].append({"instr": instr, "result": f["result"]})


def read_paragraph(p_el):
    r = _ParaReader()
    r.walk(p_el)
    while r.fields:  # 閉じていないフィールド
        r._close(r.fields.pop())
    return "".join(r.out), r.notes


def plain_ptext(p_el):
    """python-docx の p.text と同じ読み方（比較用）。"""
    return docx.text.paragraph.Paragraph(p_el, None).text


def read_document(path):
    """本文の全段落を上から順に返す。

    返り値の各要素:
      para: 通し番号（表の中の段落も数える。0から）
      pd_index: python-docx の doc.paragraphs での番号（表の中は None）
      in_table: 表の中なら [表番号, 行, 列]
      text: 抜けなく読んだ文字 / ptext: p.text の文字（比較用）
      notes: EQ・ルビなどの記録
    """
    d = docx.Document(path)
    body = d.element.body
    out = []
    pd_i = 0
    tbl_i = 0

    def add_p(p_el, in_table):
        nonlocal pd_i
        text, notes = read_paragraph(p_el)
        rec = {"para": len(out), "pd_index": None if in_table else pd_i,
               "in_table": in_table, "text": text, "ptext": plain_ptext(p_el)}
        n = {k: v for k, v in notes.items() if v}
        if n:
            rec["notes"] = n
        out.append(rec)
        if not in_table:
            pd_i += 1

    def walk_block(el, in_table=None):
        nonlocal tbl_i
        for ch in el:
            if ch.tag == W_P:
                add_p(ch, in_table)
            elif ch.tag == W_TBL:
                ti = tbl_i
                tbl_i += 1
                for ri, tr in enumerate(ch.findall(W_TR)):
                    for ci, tc in enumerate(tr.findall(W_TC)):
                        walk_block(tc, [ti, ri, ci])
            elif ch.tag in (W_SDT, W_SDTCONTENT):
                walk_block(ch, in_table)
    walk_block(body)
    return out, len(d.paragraphs), tbl_i
